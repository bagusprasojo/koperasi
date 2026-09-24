from decimal import Decimal
from django.contrib.auth import get_user_model
from django.contrib.auth.models import AnonymousUser, Group
from django.contrib.messages.storage.fallback import FallbackStorage
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.test import RequestFactory, TestCase

from core.constants import Role, CASHIER_ROLES, STAFF_ROLES
from core.context_processors import role_flags
from core.decorators import role_required
from core.rbac import get_user_roles, has_any_role, has_permission_or_role
from members.models import Member, MemberTopUp
from members.permissions import can_access_topup_proof, can_access_member_profile
from sales.models import Sale
from sales.permissions import can_view_sale_detail

User = get_user_model()


# Sample decorated views for testing
@role_required(Role.ADMIN_TOKO)
def sample_admin_view(request):
    return HttpResponse("admin ok")


@role_required(*CASHIER_ROLES)
def sample_cashier_view(request):
    return HttpResponse("cashier ok")


@role_required(Role.ADMIN_TOKO, perm='void_sale')
def sample_void_view(request):
    return HttpResponse("void ok")


class RBACTests(TestCase):
    def setUp(self):
        self.factory = RequestFactory()

        # Groups
        self.admin_group, _ = Group.objects.get_or_create(name=Role.ADMIN_TOKO)
        self.kasir_group, _ = Group.objects.get_or_create(name=Role.KASIR)
        self.pembelian_group, _ = Group.objects.get_or_create(name=Role.PEMBELIAN)
        self.member_group, _ = Group.objects.get_or_create(name=Role.MEMBER)

        # Users
        self.superuser = User.objects.create_superuser(username='superadmin', password='password123')
        self.admin_user = User.objects.create_user(username='admintoko', password='password123')
        self.admin_user.groups.add(self.admin_group)

        self.kasir_user = User.objects.create_user(username='kasir1', password='password123')
        self.kasir_user.groups.add(self.kasir_group)

        self.member_user1 = User.objects.create_user(username='member1', password='password123')
        self.member_user1.groups.add(self.member_group)
        self.member1 = Member.objects.create(user=self.member_user1, full_name='Member Satu', phone='0811111111')

        self.member_user2 = User.objects.create_user(username='member2', password='password123')
        self.member_user2.groups.add(self.member_group)
        self.member2 = Member.objects.create(user=self.member_user2, full_name='Member Dua', phone='0822222222')

    def test_superuser_bypass_without_group(self):
        """Superuser tanpa grup apapun harus diizinkan mengakses view yang diproteksi."""
        request = self.factory.get('/dummy-admin/')
        request.user = self.superuser
        response = sample_admin_view(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode(), "admin ok")

        response_cashier = sample_cashier_view(request)
        self.assertEqual(response_cashier.status_code, 200)

    def test_role_required_allows_matching_role(self):
        """User dengan role yang sesuai diizinkan mengakses view."""
        request = self.factory.get('/dummy-cashier/')
        request.user = self.kasir_user
        response = sample_cashier_view(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode(), "cashier ok")

    def test_role_required_denies_mismatched_role(self):
        """User tanpa role yang sesuai harus dilempar 403 PermissionDenied."""
        request = self.factory.get('/dummy-admin/')
        request.user = self.kasir_user  # kasir tidak punya role admin_toko
        with self.assertRaises(PermissionDenied):
            sample_admin_view(request)

    def test_role_caching_on_request(self):
        """Fungsi get_user_roles menyimpan cache di objek request."""
        request = self.factory.get('/')
        request.user = self.kasir_user

        roles_first = get_user_roles(request.user, request=request)
        self.assertIn(Role.KASIR, roles_first)
        self.assertTrue(hasattr(request, '_cached_user_roles'))

        # Query ulang melalui request harus membaca dari cache
        roles_second = get_user_roles(request.user, request=request)
        self.assertIs(roles_first, roles_second)

    def test_has_any_role_helper(self):
        self.assertTrue(has_any_role(self.superuser, Role.ADMIN_TOKO))
        self.assertTrue(has_any_role(self.kasir_user, *CASHIER_ROLES))
        self.assertFalse(has_any_role(self.kasir_user, Role.PEMBELIAN))
        self.assertFalse(has_any_role(AnonymousUser(), Role.MEMBER))

    def test_context_processor_role_flags(self):
        # Anonymous
        req_anon = self.factory.get('/')
        req_anon.user = AnonymousUser()
        flags_anon = role_flags(req_anon)
        self.assertFalse(flags_anon['is_staff_any'])
        self.assertFalse(flags_anon['is_admin_toko_role'])

        # Superuser
        req_super = self.factory.get('/')
        req_super.user = self.superuser
        flags_super = role_flags(req_super)
        self.assertTrue(flags_super['is_admin_toko_role'])
        self.assertTrue(flags_super['is_staff_any'])

        # Kasir
        req_kasir = self.factory.get('/')
        req_kasir.user = self.kasir_user
        flags_kasir = role_flags(req_kasir)
        self.assertTrue(flags_kasir['is_kasir_role'])
        self.assertTrue(flags_kasir['is_staff_any'])
        self.assertFalse(flags_kasir['is_admin_toko_role'])

        # Member
        req_member = self.factory.get('/')
        req_member.user = self.member_user1
        flags_member = role_flags(req_member)
        self.assertTrue(flags_member['is_member_role'])
        self.assertFalse(flags_member['is_staff_any'])

    def test_can_access_topup_proof(self):
        topup1 = MemberTopUp.objects.create(
            member=self.member1,
            amount=Decimal('50000.00'),
            requested_by=self.member_user1,
        )

        # Pemilik berhak
        self.assertTrue(can_access_topup_proof(self.member_user1, topup1))
        # Admin toko & Superuser berhak
        self.assertTrue(can_access_topup_proof(self.admin_user, topup1))
        self.assertTrue(can_access_topup_proof(self.superuser, topup1))
        # Member lain tidak berhak
        self.assertFalse(can_access_topup_proof(self.member_user2, topup1))
        # Kasir biasa tidak berhak jika bukan admin
        self.assertFalse(can_access_topup_proof(self.kasir_user, topup1))

    def test_can_view_sale_detail(self):
        sale1 = Sale.objects.create(
            sale_number='SL-TEST-001',
            member=self.member1,
            subtotal=Decimal('10000.00'),
            total=Decimal('10000.00'),
            created_by=self.kasir_user,
        )

        # Kasir dan Admin Toko berhak
        self.assertTrue(can_view_sale_detail(self.kasir_user, sale1))
        self.assertTrue(can_view_sale_detail(self.admin_user, sale1))
        self.assertTrue(can_view_sale_detail(self.superuser, sale1))
        # Member pemilik berhak
        self.assertTrue(can_view_sale_detail(self.member_user1, sale1))
        # Member lain tidak berhak
        self.assertFalse(can_view_sale_detail(self.member_user2, sale1))

    # --- Pengujian Pendekatan A: Manajemen Staf ---
    def test_staff_management_access_control(self):
        from core.views_staff import staff_list

        # Kasir ditolak saat mencoba membuka daftar staf
        req = self.factory.get('/staff/')
        req.user = self.kasir_user
        with self.assertRaises(PermissionDenied):
            staff_list(req)

        # Admin toko diizinkan
        req_admin = self.factory.get('/staff/')
        req_admin.user = self.admin_user
        setattr(req_admin, 'session', {})
        setattr(req_admin, '_messages', FallbackStorage(req_admin))
        resp = staff_list(req_admin)
        self.assertEqual(resp.status_code, 200)

    def test_staff_create_success(self):
        from core.views_staff import staff_create

        req = self.factory.post('/staff/create/', {
            'username': 'kasir_baru',
            'first_name': 'Budi',
            'last_name': 'Kasir',
            'email': 'budi@test.id',
            'password': 'StrongPassword123!',
            'password_confirm': 'StrongPassword123!',
            'roles': [Role.KASIR.value],
            'is_active': 'on',
        })
        req.user = self.admin_user
        setattr(req, 'session', {})
        setattr(req, '_messages', FallbackStorage(req))

        resp = staff_create(req)
        self.assertEqual(resp.status_code, 302)

        new_user = User.objects.filter(username='kasir_baru').first()
        self.assertIsNotNone(new_user)
        self.assertTrue(new_user.groups.filter(name=Role.KASIR).exists())
        self.assertTrue(new_user.is_active)

    def test_staff_edit_self_lockout_protection(self):
        from core.views_staff import staff_edit

        # Admin mencoba menonaktifkan dirinya sendiri
        req = self.factory.post(f'/staff/{self.admin_user.id}/edit/', {
            'first_name': 'Admin',
            'last_name': 'Toko',
            'email': 'admin@test.id',
            'roles': [Role.ADMIN_TOKO.value],
            # is_active tidak dicentang
        })
        req.user = self.admin_user
        setattr(req, 'session', {})
        setattr(req, '_messages', FallbackStorage(req))

        resp = staff_edit(req, self.admin_user.id)
        self.admin_user.refresh_from_db()
        self.assertTrue(self.admin_user.is_active)  # Harus tetap aktif (self-lockout dicegah)

    # --- Pengujian Pendekatan B: Matriks Izin Dinamis ---
    def test_dynamic_permission_matrix_and_enforcement(self):
        from core.views_matrix import role_permission_matrix

        # Awalnya kasir tidak punya izin void_sale
        self.assertFalse(has_permission_or_role(self.kasir_user, 'void_sale'))

        # Kasir mencoba akses view void -> Ditolak (403)
        req_void = self.factory.get('/dummy-void/')
        req_void.user = self.kasir_user
        with self.assertRaises(PermissionDenied):
            sample_void_view(req_void)

        # Admin Toko memperbarui matriks: memberikan izin void_sale ke grup Kasir
        req = self.factory.post('/staff/roles/matrix/', {
            f"perm_{Role.KASIR.value}_access_pos": 'on',
            f"perm_{Role.KASIR.value}_void_sale": 'on',
        })
        req.user = self.admin_user
        setattr(req, 'session', {})
        setattr(req, '_messages', FallbackStorage(req))

        resp = role_permission_matrix(req)
        self.assertEqual(resp.status_code, 302)

        # Reload kasir user dari DB untuk membersihkan cache permissions Django
        self.kasir_user = User.objects.get(id=self.kasir_user.id)

        # Sekarang kasir secara dinamis memiliki izin void_sale!
        self.assertTrue(has_permission_or_role(self.kasir_user, 'void_sale'))

        # Kasir sekarang berhasil mengakses view void!
        req_void2 = self.factory.get('/dummy-void/')
        req_void2.user = self.kasir_user
        resp_void = sample_void_view(req_void2)
        self.assertEqual(resp_void.status_code, 200)
        self.assertEqual(resp_void.content.decode(), "void ok")

    def test_menu_dynamic_permissions_rendering(self):
        from django.contrib.auth.models import Permission
        core_perms = {p.codename: p for p in Permission.objects.filter(content_type__app_label='core')}
        if 'access_pos' in core_perms:
            self.kasir_group.permissions.set([
                core_perms['access_pos'],
                core_perms['view_inventory'],
                core_perms['view_members'],
                core_perms['view_sales_reports'],
            ])
        if 'manage_purchases' in core_perms:
            self.pembelian_group.permissions.set([
                core_perms['manage_purchases'],
                core_perms['manage_products'],
                core_perms['view_inventory'],
                core_perms['view_members'],
                core_perms['view_sales_reports'],
            ])
        if 'manage_staff' in core_perms:
            self.admin_group.permissions.set(list(core_perms.values()))

        pembelian_user = User.objects.create_user(username='pembelian1', password='password123')
        pembelian_user.groups.add(self.pembelian_group)

        # 1. Kasir: Melihat Kasir (POS), tetapi TIDAK melihat Kulakan atau Kelola Staf
        self.client.force_login(self.kasir_user)
        resp_kasir = self.client.get('/dashboard/')
        self.assertEqual(resp_kasir.status_code, 200)
        self.assertContains(resp_kasir, 'Kasir (POS)')
        self.assertNotContains(resp_kasir, 'Kulakan')
        self.assertNotContains(resp_kasir, 'Kelola Staf')

        # 2. Pembelian: Melihat Kulakan, tetapi TIDAK melihat Kasir (POS) atau Kelola Staf
        self.client.force_login(pembelian_user)
        resp_pembelian = self.client.get('/dashboard/')
        self.assertEqual(resp_pembelian.status_code, 200)
        self.assertContains(resp_pembelian, 'Kulakan')
        self.assertNotContains(resp_pembelian, 'Kasir (POS)')
        self.assertNotContains(resp_pembelian, 'Kelola Staf')

        # 3. Admin Toko: Melihat seluruh menu (Kasir POS, Kulakan, Kelola Staf)
        self.client.force_login(self.admin_user)
        resp_admin = self.client.get('/dashboard/')
        self.assertEqual(resp_admin.status_code, 200)
        self.assertContains(resp_admin, 'Kasir (POS)')
        self.assertContains(resp_admin, 'Kulakan')
        self.assertContains(resp_admin, 'Kelola Staf')

        # 4. Member: Melihat portal member, TIDAK melihat modul staf
        self.client.force_login(self.member_user1)
        resp_member = self.client.get('/dashboard/')
        self.assertEqual(resp_member.status_code, 200)
        self.assertContains(resp_member, 'Topup Saya')
        self.assertNotContains(resp_member, 'Kasir (POS)')
        self.assertNotContains(resp_member, 'Kulakan')
        self.assertNotContains(resp_member, 'Kelola Staf')

    def test_insufficient_stock_error_message_detailed(self):
        from inventory.models import Category, Unit, Product, ProductPriceTier
        from inventory.services import post_pos_sale
        from django.core.exceptions import ValidationError

        cat = Category.objects.create(name='Sembako')
        unit = Unit.objects.create(name='Botol', code='BTL')
        prod = Product.objects.create(
            category=cat,
            unit=unit,
            name='Kecap Manis 600ml',
            sku='KCP-600',
            stock=2,
            cost_of_goods_sold=Decimal('15000.00'),
        )

        with self.assertRaises(ValidationError) as ctx:
            post_pos_sale(product=prod, qty=5, user=self.admin_user)

        err_msg = str(ctx.exception)
        self.assertIn("Kecap Manis 600ml", err_msg)
        self.assertIn("Tersedia: 2 Botol", err_msg)
        self.assertIn("diminta: 5 Botol", err_msg)

    def test_checkout_pos_pre_validation_lists_all_insufficient_items(self):
        from inventory.models import Category, Unit, Product, ProductPriceTier
        from sales.models import SalePayment
        from sales.services import checkout_pos
        from django.core.exceptions import ValidationError

        cat = Category.objects.create(name='Minuman')
        unit = Unit.objects.create(name='Pcs', code='PCS')
        p1 = Product.objects.create(
            category=cat,
            unit=unit,
            name='Kopi Hitam',
            sku='KOP-01',
            stock=1,
            cost_of_goods_sold=Decimal('3000.00'),
        )
        ProductPriceTier.objects.create(product=p1, level=1, min_qty=1, max_qty=999, price=Decimal('5000.00'))

        p2 = Product.objects.create(
            category=cat,
            unit=unit,
            name='Teh Celup',
            sku='TEH-01',
            stock=0,
            cost_of_goods_sold=Decimal('2000.00'),
        )
        ProductPriceTier.objects.create(product=p2, level=1, min_qty=1, max_qty=999, price=Decimal('4000.00'))

        items = [
            {'product_id': str(p1.id), 'qty': 3},
            {'product_id': str(p2.id), 'qty': 2},
        ]
        payments = [{'method': SalePayment.METHOD_CASH, 'amount': '23000.00'}]

        with self.assertRaises(ValidationError) as ctx:
            checkout_pos(
                member_id=None,
                items=items,
                payments=payments,
                client_txn_id='test-stock-val-1',
                user=self.admin_user,
                cash_received_raw='23000.00',
            )

        err_msg = str(ctx.exception)
        # Kedua produk yang kurang stoknya harus tercantum dalam satu pesan error
        self.assertIn("Kopi Hitam", err_msg)
        self.assertIn("sisa stok 1 Pcs, diminta 3 Pcs", err_msg)
        self.assertIn("Teh Celup", err_msg)
        self.assertIn("sisa stok 0 Pcs, diminta 2 Pcs", err_msg)


class ProductImportExcelTests(TestCase):
    def setUp(self):
        import openpyxl
        from django.contrib.auth import get_user_model
        from django.contrib.auth.models import Group
        from core.constants import Role
        from inventory.models import Category, Unit, Product

        User = get_user_model()
        self.admin_group, _ = Group.objects.get_or_create(name=Role.ADMIN_TOKO)
        self.admin_user = User.objects.create_user(username='admin_import', password='password123')
        self.admin_user.groups.add(self.admin_group)

        self.cat_sembako = Category.objects.create(name='Sembako')
        self.unit_pcs = Unit.objects.create(name='Pieces', code='PCS', is_active=True)
        self.unit_kg = Unit.objects.create(name='Kilogram', code='KG', is_active=True)

    def _make_excel_file(self, rows, sheet_name='Template Import'):
        import io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Headers
        ws.append(['sku', 'barcode', 'nama_barang', 'kategori', 'satuan', 'harga_beli', 'harga_jual', 'min_stok', 'stok_awal'])
        for r in rows:
            ws.append(r)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return SimpleUploadedFile('test_import.xlsx', out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_download_template_excel(self):
        import openpyxl
        import io
        self.client.force_login(self.admin_user)
        resp = self.client.get('/inventory/products/import/template/')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', resp['Content-Type'])
        self.assertIn('template_import_produk.xlsx', resp['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn('Template Import', wb.sheetnames)
        self.assertIn('Referensi Master Data', wb.sheetnames)

    def test_preview_rejects_missing_category_no_autocreate(self):
        from inventory.models import Category
        self.client.force_login(self.admin_user)

        rows = [
            ['BRG-99', '', 'Barang Aneh', 'KategoriTidakAda', 'PCS', 10000, 15000, 5, 10]
        ]
        excel_file = self._make_excel_file(rows)
        resp = self.client.post('/inventory/products/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp.status_code, 200)

        # Harus masuk mode preview
        self.assertTrue(resp.context['preview_mode'])
        self.assertFalse(resp.context['can_confirm'])
        self.assertEqual(resp.context['total_errors'], 1)

        err_list = resp.context['rows'][0]['errors']
        self.assertTrue(any("Kategori 'KategoriTidakAda' tidak ditemukan di data master" in e for e in err_list))

        # Pastikan TIDAK di-auto-create
        self.assertFalse(Category.objects.filter(name='KategoriTidakAda').exists())

    def test_preview_rejects_missing_unit_no_autocreate(self):
        from inventory.models import Unit
        self.client.force_login(self.admin_user)

        rows = [
            ['BRG-98', '', 'Minyak Goreng', 'Sembako', 'LUSIN_GAIB', 10000, 15000, 5, 10]
        ]
        excel_file = self._make_excel_file(rows)
        resp = self.client.post('/inventory/products/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp.status_code, 200)

        self.assertTrue(resp.context['preview_mode'])
        self.assertFalse(resp.context['can_confirm'])
        self.assertEqual(resp.context['total_errors'], 1)

        err_list = resp.context['rows'][0]['errors']
        self.assertTrue(any("Satuan 'LUSIN_GAIB' tidak ditemukan atau nonaktif di data master" in e for e in err_list))

        # Pastikan TIDAK di-auto-create
        self.assertFalse(Unit.objects.filter(code='LUSIN_GAIB').exists())

    def test_preview_rejects_existing_sku_in_database(self):
        from inventory.models import Product
        # Buat produk di DB
        Product.objects.create(
            category=self.cat_sembako,
            unit=self.unit_pcs,
            name='Produk Lama',
            sku='SKU-EXIST-01',
            stock=5,
        )

        self.client.force_login(self.admin_user)
        rows = [
            ['SKU-EXIST-01', '', 'Produk Duplikat DB', 'Sembako', 'PCS', 5000, 8000, 2, 10]
        ]
        excel_file = self._make_excel_file(rows)
        resp = self.client.post('/inventory/products/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp.status_code, 200)

        self.assertFalse(resp.context['can_confirm'])
        self.assertEqual(resp.context['total_errors'], 1)

        err_list = resp.context['rows'][0]['errors']
        self.assertTrue(any("SKU 'SKU-EXIST-01' sudah terdaftar di database" in e for e in err_list))

    def test_preview_rejects_duplicate_sku_inside_excel_file(self):
        self.client.force_login(self.admin_user)
        rows = [
            ['SKU-SAME-01', '', 'Produk A', 'Sembako', 'PCS', 5000, 8000, 2, 10],
            ['SKU-SAME-01', '', 'Produk B', 'Sembako', 'PCS', 6000, 9000, 2, 10],
        ]
        excel_file = self._make_excel_file(rows)
        resp = self.client.post('/inventory/products/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp.status_code, 200)

        self.assertFalse(resp.context['can_confirm'])
        self.assertEqual(resp.context['total_errors'], 1)

        # Baris ke-2 harus menandai duplikasi internal
        err_list = resp.context['rows'][1]['errors']
        self.assertTrue(any("duplikat dengan baris 2 di file Excel" in e for e in err_list))

    def test_confirm_blocked_when_session_batch_has_errors(self):
        from inventory.models import Product
        self.client.force_login(self.admin_user)

        # Set fake session dengan can_confirm = False
        session = self.client.session
        session['product_import_batch'] = {
            'rows': [{'sku': 'SKU-X'}],
            'can_confirm': False,
            'total_rows': 1,
            'total_errors': 1,
            'total_valid': 0,
        }
        session.save()

        resp = self.client.post('/inventory/products/import/', {'action': 'confirm'}, follow=True)
        self.assertEqual(resp.status_code, 200)
        # Harus tolak simpan
        self.assertFalse(Product.objects.filter(sku='SKU-X').exists())

    def test_successful_import_and_atomic_stock_ledger(self):
        from inventory.models import Product, ProductPriceTier, StockLedger
        self.client.force_login(self.admin_user)

        rows = [
            ['SKU-OK-01', '899111', 'Beras Super 5kg', 'Sembako', 'PCS', 50000, 60000, 5, 25],
            ['SKU-OK-02', '', 'Tepung Terigu 1kg', 'Sembako', 'KG', 10000, 12000, 10, 0],
        ]
        excel_file = self._make_excel_file(rows)

        # 1. Preview
        resp_preview = self.client.post('/inventory/products/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp_preview.status_code, 200)
        self.assertTrue(resp_preview.context['can_confirm'])
        self.assertEqual(resp_preview.context['total_errors'], 0)
        self.assertEqual(resp_preview.context['total_valid'], 2)

        # 2. Confirm
        resp_confirm = self.client.post('/inventory/products/import/', {'action': 'confirm'}, follow=True)
        self.assertEqual(resp_confirm.status_code, 200)

        # Verifikasi produk 1 (dengan saldo awal 25)
        p1 = Product.objects.get(sku='SKU-OK-01')
        self.assertEqual(p1.name, 'Beras Super 5kg')
        self.assertEqual(p1.category, self.cat_sembako)
        self.assertEqual(p1.unit, self.unit_pcs)
        self.assertEqual(p1.stock, 25)
        self.assertEqual(p1.reorder_point, 5)

        # Verifikasi tier harga
        tier1 = ProductPriceTier.objects.get(product=p1, level=1)
        self.assertEqual(tier1.price, Decimal('60000.00'))

        # Verifikasi kartu stok
        ledger1 = StockLedger.objects.filter(product=p1).first()
        self.assertIsNotNone(ledger1)
        self.assertEqual(ledger1.qty_in, 25)
        self.assertEqual(ledger1.balance_after, 25)

        # Verifikasi produk 2 (stok awal 0)
        p2 = Product.objects.get(sku='SKU-OK-02')
        self.assertEqual(p2.name, 'Tepung Terigu 1kg')
        self.assertEqual(p2.stock, 0)
        self.assertEqual(p2.reorder_point, 10)
        self.assertFalse(StockLedger.objects.filter(product=p2).exists())


class MemberImportExcelTests(TestCase):
    def setUp(self):
        self.admin_group, _ = Group.objects.get_or_create(name=Role.ADMIN_TOKO)
        self.admin_user = User.objects.create_user(username='admin_import_mbr', password='password123')
        self.admin_user.groups.add(self.admin_group)

    def _make_excel_file(self, rows, sheet_name='Template Import Member'):
        import io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Headers
        ws.append(['kode_member', 'nama_lengkap', 'telepon', 'email', 'alamat', 'nomor_kartu', 'saldo_awal', 'password'])
        for r in rows:
            ws.append(r)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return SimpleUploadedFile('test_members_import.xlsx', out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_download_template_excel(self):
        import openpyxl
        import io
        self.client.force_login(self.admin_user)
        resp = self.client.get('/members/import/template/')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', resp['Content-Type'])
        self.assertIn('template_import_member.xlsx', resp['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn('Template Import Member', wb.sheetnames)
        self.assertIn('Panduan Pengisian', wb.sheetnames)

    def test_preview_rejects_duplicate_code_in_db(self):
        existing_user = User.objects.create_user(username='mbr_dummy', password='password123')
        Member.objects.create(code='MBR-EXIST', user=existing_user, full_name='Dummy Member', phone='0811119999')

        self.client.force_login(self.admin_user)
        rows = [
            ['MBR-EXIST', 'Budi Dua', '0899999999', 'budi@test.com', 'Alamat', '', 0, '']
        ]
        excel_file = self._make_excel_file(rows)
        resp = self.client.post('/members/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.context['preview_mode'])
        self.assertFalse(resp.context['can_confirm'])
        self.assertEqual(resp.context['total_errors'], 1)
        err_list = resp.context['rows'][0]['errors']
        self.assertTrue(any("sudah terdaftar" in e for e in err_list))

    def test_preview_rejects_duplicate_phone_in_db(self):
        existing_user = User.objects.create_user(username='mbr_dummy2', password='password123')
        Member.objects.create(code='MBR-DUMMY2', user=existing_user, full_name='Dummy Member 2', phone='08123456789')

        self.client.force_login(self.admin_user)
        rows = [
            ['MBR-BARU', 'Budi Baru', '08123456789', 'budi@test.com', 'Alamat', '', 0, '']
        ]
        excel_file = self._make_excel_file(rows)
        resp = self.client.post('/members/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.context['preview_mode'])
        self.assertFalse(resp.context['can_confirm'])
        self.assertEqual(resp.context['total_errors'], 1)
        err_list = resp.context['rows'][0]['errors']
        self.assertTrue(any("sudah terdaftar" in e for e in err_list))

    def test_preview_rejects_duplicate_card_in_db(self):
        from members.models import MemberCard
        existing_user = User.objects.create_user(username='mbr_dummy3', password='password123')
        m3 = Member.objects.create(code='MBR-DUMMY3', user=existing_user, full_name='Dummy Member 3', phone='08123456780')
        MemberCard.objects.create(member=m3, card_number='CRD-EXIST', status=MemberCard.STATUS_ACTIVE)

        self.client.force_login(self.admin_user)
        rows = [
            ['MBR-BARU3', 'Budi Baru 3', '0899999988', 'budi@test.com', 'Alamat', 'CRD-EXIST', 0, '']
        ]
        excel_file = self._make_excel_file(rows)
        resp = self.client.post('/members/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.context['preview_mode'])
        self.assertFalse(resp.context['can_confirm'])
        self.assertEqual(resp.context['total_errors'], 1)
        err_list = resp.context['rows'][0]['errors']
        self.assertTrue(any("sudah terdaftar" in e for e in err_list))

    def test_preview_rejects_internal_duplicates_in_file(self):
        self.client.force_login(self.admin_user)
        rows = [
            ['MBR-DUP', 'Member A', '0812000001', '', '', '', 0, ''],
            ['MBR-DUP', 'Member B', '0812000002', '', '', '', 0, ''],
        ]
        excel_file = self._make_excel_file(rows)
        resp = self.client.post('/members/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(resp.context['can_confirm'])
        self.assertGreaterEqual(resp.context['total_errors'], 1)

    def test_confirm_blocked_when_session_batch_has_errors(self):
        self.client.force_login(self.admin_user)
        session = self.client.session
        session['member_import_batch'] = {
            'rows': [{'kode_member': 'MBR-ERR', 'nama_lengkap': 'Error Mbr', 'telepon': '081', 'is_valid': False, 'errors': ['Error']}],
            'can_confirm': False,
            'total_rows': 1,
            'total_errors': 1,
            'total_valid': 0,
        }
        session.save()

        resp = self.client.post('/members/import/', {'action': 'confirm'}, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(Member.objects.filter(code='MBR-ERR').exists())

    def test_successful_import_and_atomic_entities(self):
        from members.models import MemberCard, MemberWallet, MemberLedger, MemberDepositAuditLog
        self.client.force_login(self.admin_user)

        rows = [
            ['MBR-OK-01', 'Budi Santoso', '0899123401', 'budi@example.com', 'Sleman', 'CRD-901', 50000, 'rahasia123'],
            ['MBR-OK-02', 'Siti Aminah', '0899123402', '', 'Bantul', '', 0, ''],
        ]
        excel_file = self._make_excel_file(rows)

        # 1. Preview
        resp_preview = self.client.post('/members/import/', {'action': 'preview', 'file': excel_file})
        self.assertEqual(resp_preview.status_code, 200)
        self.assertTrue(resp_preview.context['can_confirm'])
        self.assertEqual(resp_preview.context['total_errors'], 0)
        self.assertEqual(resp_preview.context['total_valid'], 2)

        # 2. Confirm
        resp_confirm = self.client.post('/members/import/', {'action': 'confirm'}, follow=True)
        self.assertEqual(resp_confirm.status_code, 200)

        # Verifikasi Member 1 (dengan saldo awal 50000 dan kartu custom)
        m1 = Member.objects.get(code='MBR-OK-01')
        self.assertEqual(m1.full_name, 'Budi Santoso')
        self.assertEqual(m1.phone, '0899123401')
        self.assertEqual(m1.email, 'budi@example.com')
        self.assertEqual(m1.address, 'Sleman')
        self.assertIsNotNone(m1.user)
        self.assertEqual(m1.user.username, 'MBR-OK-01')
        self.assertTrue(m1.user.check_password('rahasia123'))
        self.assertTrue(m1.user.groups.filter(name=Role.MEMBER).exists())

        c1 = MemberCard.objects.get(member=m1)
        self.assertEqual(c1.card_number, 'CRD-901')
        self.assertEqual(c1.status, MemberCard.STATUS_ACTIVE)

        w1 = MemberWallet.objects.get(member=m1)
        self.assertEqual(w1.balance, Decimal('50000.00'))

        ledger1 = MemberLedger.objects.filter(member=m1).first()
        self.assertIsNotNone(ledger1)
        self.assertEqual(ledger1.txn_type, MemberLedger.TYPE_TOPUP)
        self.assertEqual(ledger1.amount, Decimal('50000.00'))
        self.assertEqual(ledger1.balance_after, Decimal('50000.00'))

        audit1 = MemberDepositAuditLog.objects.filter(member=m1).first()
        self.assertIsNotNone(audit1)
        self.assertEqual(audit1.amount, Decimal('50000.00'))

        # Verifikasi Member 2 (saldo awal 0, default kartu = kode, default password = telepon)
        m2 = Member.objects.get(code='MBR-OK-02')
        self.assertEqual(m2.full_name, 'Siti Aminah')
        self.assertEqual(m2.phone, '0899123402')
        self.assertIsNotNone(m2.user)
        self.assertEqual(m2.user.username, 'MBR-OK-02')
        self.assertTrue(m2.user.check_password('0899123402'))

        c2 = MemberCard.objects.get(member=m2)
        self.assertEqual(c2.card_number, 'MBR-OK-02')

        w2 = MemberWallet.objects.get(member=m2)
        self.assertEqual(w2.balance, Decimal('0.00'))
        self.assertFalse(MemberLedger.objects.filter(member=m2).exists())



