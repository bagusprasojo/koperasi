from django.core.paginator import Paginator
from django.contrib import messages
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.http import url_has_allowed_host_and_scheme
from django.db.models import Q
from django.db import IntegrityError, transaction
from django.db.models.deletion import ProtectedError
from django.core.exceptions import ValidationError
from decimal import Decimal, InvalidOperation
import json
from datetime import date, timedelta

from core.decorators import role_required
from core.constants import Role, STAFF_ROLES, MANAGEMENT_ROLES
from .models import Category, DailyClosing, InventoryTransaction, InventoryTransactionItem, Product, ProductPriceTier, Supplier, Unit
from .services import (
    close_daily,
    create_purchase_transaction,
    delete_purchase_transaction,
    edit_purchase_transaction,
    post_internal_used,
    post_purchase,
    post_stock_opname,
    reopen_last_closing,
    low_stock_products,
    import_products_batch,
)

TWOPLACES = Decimal('0.01')


def _to_json_payload(value):
    return json.dumps(value, default=str)


def _exc_message(exc):
    if isinstance(exc, ValidationError):
        return exc.messages[0] if exc.messages else str(exc)
    return str(exc)


def _get_safe_next_url(request):
    next_url = request.GET.get('next') or request.POST.get('next') or ''
    if not url_has_allowed_host_and_scheme(
        url=next_url,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return ''
    return next_url


def _extract_tier_rows(request):
    rows = []
    for level in [1, 2, 3]:
        rows.append(
            {
                'level': str(level),
                'min_qty': request.POST.get(f'tier_{level}_min_qty', '').strip(),
                'max_qty': request.POST.get(f'tier_{level}_max_qty', '').strip(),
                'mode': request.POST.get(f'tier_{level}_mode', 'final').strip(),
                'discount_type': request.POST.get(f'tier_{level}_discount_type', 'percent').strip(),
                'input_value': request.POST.get(f'tier_{level}_value', '').strip(),
                'price': '',
            }
        )
    return rows


def _serialize_tiers(product):
    by_level = {t.level: t for t in product.price_tiers.order_by('level')}
    rows = []
    for level in [1, 2, 3]:
        tier = by_level.get(level)
        if tier:
            rows.append(
                {
                    'level': str(level),
                    'min_qty': str(tier.min_qty),
                    'max_qty': str(tier.max_qty),
                    'mode': tier.source_mode or 'final',
                    'discount_type': tier.discount_type or 'percent',
                    'input_value': str(tier.discount_value if tier.source_mode == 'discount' and tier.discount_value is not None else tier.price),
                    'price': str(tier.price),
                }
            )
        else:
            rows.append(
                {
                    'level': str(level),
                    'min_qty': '',
                    'max_qty': '',
                    'mode': 'final',
                    'discount_type': 'percent',
                    'input_value': '',
                    'price': '',
                }
            )
    return rows


def _validate_contiguous_ranges(rows):
    prev_max = None
    for idx, row in enumerate(rows):
        min_qty = int(row['min_qty'])
        max_qty = int(row['max_qty'])
        if min_qty > max_qty:
            raise ValidationError(f'Level {idx + 1}: min qty tidak boleh lebih besar dari max qty.')
        if idx == 0 and min_qty != 1:
            raise ValidationError('Level 1 harus dimulai dari qty 1.')
        if prev_max is not None and min_qty != (prev_max + 1):
            raise ValidationError(
                f'Range level {idx} ke level {idx + 1} harus berurutan tanpa gap/overlap.'
            )
        prev_max = max_qty


def _compute_level_prices(rows):
    # Level aktif ditentukan dari pengisian salah satu field inti.
    active_flags = [
        bool(rows[i]['min_qty'] or rows[i]['max_qty'] or rows[i]['input_value'])
        for i in range(3)
    ]

    if not active_flags[0]:
        raise ValidationError('Minimal level 1 wajib diisi.')

    # Tidak boleh lompat level: contoh level 1 + 3 tanpa level 2.
    found_inactive = False
    for i, is_active in enumerate(active_flags):
        if not is_active:
            found_inactive = True
            continue
        if found_inactive:
            raise ValidationError(f'Level {i + 1} tidak boleh diisi jika level sebelumnya kosong.')

    active_rows = rows[:sum(1 for flag in active_flags if flag)]
    for idx, row in enumerate(active_rows):
        if not all([row['min_qty'], row['max_qty'], row['input_value']]):
            raise ValidationError(f'Semua field range dan nilai harga pada level {idx + 1} wajib diisi.')

    _validate_contiguous_ranges(active_rows)

    base_price = Decimal(active_rows[0]['input_value']).quantize(TWOPLACES)
    if base_price <= 0:
        raise ValidationError('Level 1 wajib memiliki harga lebih besar dari 0.')
    active_rows[0]['price'] = base_price
    active_rows[0]['source_mode'] = 'final'
    active_rows[0]['discount_type'] = ''
    active_rows[0]['discount_value'] = None

    for i in range(1, len(active_rows)):
        mode = active_rows[i]['mode']
        value = Decimal(active_rows[i]['input_value']).quantize(TWOPLACES)
        if mode == 'final':
            price = value
            active_rows[i]['source_mode'] = 'final'
            active_rows[i]['discount_type'] = ''
            active_rows[i]['discount_value'] = None
        else:
            if active_rows[i]['discount_type'] == 'percent':
                if value < 0 or value > 100:
                    raise ValidationError(f'Level {i + 1}: diskon persen harus 0 sampai 100.')
                price = (base_price - ((value / Decimal('100')) * base_price)).quantize(TWOPLACES)
                active_rows[i]['source_mode'] = 'discount'
                active_rows[i]['discount_type'] = 'percent'
                active_rows[i]['discount_value'] = value
            else:
                if value < 0:
                    raise ValidationError(f'Level {i + 1}: diskon nominal tidak boleh negatif.')
                price = (base_price - value).quantize(TWOPLACES)
                active_rows[i]['source_mode'] = 'discount'
                active_rows[i]['discount_type'] = 'nominal'
                active_rows[i]['discount_value'] = value
        if price <= 0:
            raise ValidationError(f'Level {i + 1}: harga akhir harus lebih besar dari 0.')
        active_rows[i]['price'] = price.quantize(TWOPLACES)

    for i in range(1, len(active_rows)):
        if active_rows[i - 1]['price'] < active_rows[i]['price']:
            raise ValidationError('Harga level harus menurun atau sama: Level 1 >= Level 2 >= Level 3.')

    return active_rows


def _create_price_tiers(product, tier_rows):
    computed_rows = _compute_level_prices(tier_rows)
    for row in computed_rows:
        ProductPriceTier.objects.create(
            product=product,
            level=int(row['level']),
            min_qty=int(row['min_qty']),
            max_qty=int(row['max_qty']),
            price=Decimal(row['price']),
            source_mode=row['source_mode'],
            discount_type=row['discount_type'],
            discount_value=row['discount_value'],
        )


def _parse_money(raw_value, field_label):
    value_raw = (raw_value or '').strip()
    if value_raw == '':
        return Decimal('0.00')
    try:
        value = Decimal(value_raw).quantize(TWOPLACES)
    except (InvalidOperation, ValueError):
        raise ValidationError(f'{field_label} tidak valid.')
    if value < 0:
        raise ValidationError(f'{field_label} tidak boleh negatif.')
    return value


def _parse_non_negative_int(raw_value, field_label):
    raw = (raw_value or '').strip()
    if raw == '':
        return 0
    try:
        value = int(raw)
    except (ValueError, TypeError):
        raise ValidationError(f'{field_label} harus berupa bilangan bulat.')
    if value < 0:
        raise ValidationError(f'{field_label} tidak boleh negatif.')
    return value


def _latest_supplier_for_product(product):
    item = (
        InventoryTransactionItem.objects
        .select_related('transaction__supplier')
        .filter(
            product=product,
            transaction__tx_type=InventoryTransaction.TYPE_PURCHASE,
            transaction__supplier__isnull=False,
        )
        .order_by('-transaction__tx_date', '-transaction__created_at')
        .first()
    )
    return item.transaction.supplier if item else None


def _default_tier_rows():
    return [
        {
            'level': '1',
            'min_qty': '1',
            'max_qty': '',
            'mode': 'final',
            'discount_type': 'percent',
            'input_value': '',
            'price': '',
        },
        {
            'level': '2',
            'min_qty': '',
            'max_qty': '',
            'mode': 'discount',
            'discount_type': 'percent',
            'input_value': '',
            'price': '',
        },
        {
            'level': '3',
            'min_qty': '',
            'max_qty': '',
            'mode': 'discount',
            'discount_type': 'percent',
            'input_value': '',
            'price': '',
        }
    ]


@role_required(*STAFF_ROLES, perm='view_inventory')
def product_list(request):
    query = request.GET.get('q', '').strip()
    products = Product.objects.select_related('category', 'unit').prefetch_related('price_tiers').order_by('name')
    if query:
        products = products.filter(
            Q(name__icontains=query) |
            Q(sku__icontains=query) |
            Q(barcode__icontains=query) |
            Q(category__name__icontains=query) |
            Q(unit__name__icontains=query) |
            Q(unit__code__icontains=query)
        )

    paginator = Paginator(products, 10)
    page_obj = paginator.get_page(request.GET.get('page'))
    return render(
        request,
        'inventory/product_list.html',
        {'page_obj': page_obj, 'query': query},
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def product_create(request):
    error_message = ''
    categories = Category.objects.order_by('name')
    units = Unit.objects.filter(is_active=True).order_by('name')
    tier_rows = _default_tier_rows()
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        sku = request.POST.get('sku', '').strip()
        barcode = request.POST.get('barcode', '').strip()
        category_id = request.POST.get('category_id', '').strip()
        unit_id = request.POST.get('unit_id', '').strip()
        last_purchase_price = request.POST.get('last_purchase_price', '').strip()
        cost_of_goods_sold = request.POST.get('cost_of_goods_sold', '').strip()
        reorder_point_raw = request.POST.get('reorder_point', '').strip()
        tier_rows = _extract_tier_rows(request)

        if not all([name, sku, category_id, unit_id]):
            error_message = 'Nama, SKU, kategori, dan satuan wajib diisi.'
            messages.error(request, error_message)
        else:
            try:
                category = Category.objects.get(id=category_id)
                unit = Unit.objects.get(id=unit_id, is_active=True)
                buy_price = _parse_money(last_purchase_price, 'Harga beli')
                hpp = _parse_money(cost_of_goods_sold, 'Harga pokok penjualan')
                reorder_point = _parse_non_negative_int(reorder_point_raw, 'Reorder point')
                if barcode and Product.objects.filter(barcode=barcode).exists():
                    error_message = 'Barcode sudah dipakai.'
                    messages.error(request, error_message)
                    return render(
                        request,
                        'inventory/product_create.html',
                        {
                            'error_message': error_message,
                            'categories': categories,
                            'units': units,
                            'tier_rows': tier_rows,
                            'tier_rows_json': _to_json_payload(tier_rows),
                        },
                    )
                with transaction.atomic():
                    product = Product.objects.create(
                        name=name,
                        sku=sku,
                        barcode=barcode or None,
                        stock=0,
                        last_purchase_price=buy_price,
                        cost_of_goods_sold=hpp,
                        reorder_point=reorder_point,
                        category=category,
                        unit=unit,
                        allow_decimal_qty=bool(request.POST.get('allow_decimal_qty')),
                    )
                    _create_price_tiers(product, tier_rows)
                messages.success(request, 'Produk berhasil ditambahkan.')
                return redirect('product_list')
            except (ValueError, InvalidOperation):
                error_message = 'Format level harga tidak valid.'
                messages.error(request, error_message)
            except ValidationError as exc:
                error_message = exc.messages[0] if exc.messages else str(exc)
                messages.error(request, error_message)
            except Category.DoesNotExist:
                error_message = 'Kategori tidak valid.'
                messages.error(request, error_message)
            except Unit.DoesNotExist:
                error_message = 'Satuan tidak valid atau nonaktif.'
                messages.error(request, error_message)
            except IntegrityError:
                error_message = 'SKU sudah dipakai atau data tidak valid.'
                messages.error(request, error_message)

    return render(
        request,
        'inventory/product_create.html',
        {
            'error_message': error_message,
            'categories': categories,
            'units': units,
            'latest_supplier': None,
            'tier_rows': tier_rows,
            'tier_rows_json': _to_json_payload(tier_rows),
        },
    )


@role_required(*STAFF_ROLES, perm='view_inventory')
def product_detail(request, uuid):
    product = get_object_or_404(
        Product.objects.select_related('category', 'unit').prefetch_related('price_tiers'),
        uuid=uuid,
    )
    return render(request, 'inventory/product_detail.html', {'product': product})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def product_edit(request, uuid):
    product = get_object_or_404(Product.objects.prefetch_related('price_tiers'), uuid=uuid)
    categories = Category.objects.order_by('name')
    units = Unit.objects.filter(is_active=True).order_by('name')
    error_message = ''
    next_url = _get_safe_next_url(request)
    back_url = next_url or f"/inventory/products/{product.uuid}/"
    tier_rows = _serialize_tiers(product) or _default_tier_rows()
    latest_supplier = _latest_supplier_for_product(product)

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        sku = request.POST.get('sku', '').strip()
        barcode = request.POST.get('barcode', '').strip()
        category_id = request.POST.get('category_id', '').strip()
        unit_id = request.POST.get('unit_id', '').strip()
        last_purchase_price = request.POST.get('last_purchase_price', '').strip()
        cost_of_goods_sold = request.POST.get('cost_of_goods_sold', '').strip()
        reorder_point_raw = request.POST.get('reorder_point', '').strip()
        tier_rows = _extract_tier_rows(request)

        if not all([name, sku, category_id, unit_id]):
            error_message = 'Nama, SKU, kategori, dan satuan wajib diisi.'
            messages.error(request, error_message)
        else:
            try:
                category = Category.objects.get(id=category_id)
                unit = Unit.objects.get(id=unit_id, is_active=True)
                buy_price = _parse_money(last_purchase_price, 'Harga beli')
                hpp = _parse_money(cost_of_goods_sold, 'Harga pokok penjualan')
                reorder_point = _parse_non_negative_int(reorder_point_raw, 'Reorder point')
                if Product.objects.exclude(id=product.id).filter(sku=sku).exists():
                    error_message = 'SKU sudah dipakai.'
                    messages.error(request, error_message)
                    return render(
                        request,
                        'inventory/product_edit.html',
                        {
                            'product': product,
                            'categories': categories,
                            'units': units,
                            'error_message': error_message,
                            'back_url': back_url,
                            'next_url': next_url,
                            'latest_supplier': latest_supplier,
                            'tier_rows': tier_rows,
                            'tier_rows_json': _to_json_payload(tier_rows),
                        },
                    )
                if barcode and Product.objects.exclude(id=product.id).filter(barcode=barcode).exists():
                    error_message = 'Barcode sudah dipakai.'
                    messages.error(request, error_message)
                    return render(
                        request,
                        'inventory/product_edit.html',
                        {
                            'product': product,
                            'categories': categories,
                            'units': units,
                            'error_message': error_message,
                            'back_url': back_url,
                            'next_url': next_url,
                            'latest_supplier': latest_supplier,
                            'tier_rows': tier_rows,
                            'tier_rows_json': _to_json_payload(tier_rows),
                        },
                    )
                with transaction.atomic():
                    product.name = name
                    product.sku = sku
                    product.barcode = barcode or None
                    product.category = category
                    product.unit = unit
                    product.last_purchase_price = buy_price
                    product.cost_of_goods_sold = hpp
                    product.reorder_point = reorder_point
                    product.allow_decimal_qty = bool(request.POST.get('allow_decimal_qty'))
                    product.save()
                    product.price_tiers.all().delete()
                    _create_price_tiers(product, tier_rows)
                messages.success(request, 'Produk berhasil diperbarui.')
                return redirect('product_list')
            except (ValueError, InvalidOperation):
                error_message = 'Format level harga tidak valid.'
                messages.error(request, error_message)
            except ValidationError as exc:
                error_message = exc.messages[0] if exc.messages else str(exc)
                messages.error(request, error_message)
            except Category.DoesNotExist:
                error_message = 'Kategori tidak valid.'
                messages.error(request, error_message)
            except Unit.DoesNotExist:
                error_message = 'Satuan tidak valid atau nonaktif.'
                messages.error(request, error_message)
            except IntegrityError:
                error_message = 'SKU sudah dipakai atau data tidak valid.'
                messages.error(request, error_message)

    return render(
        request,
        'inventory/product_edit.html',
        {
            'product': product,
            'categories': categories,
            'units': units,
            'error_message': error_message,
            'back_url': back_url,
            'next_url': next_url,
            'latest_supplier': latest_supplier,
            'tier_rows': tier_rows,
            'tier_rows_json': _to_json_payload(tier_rows),
        },
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def product_delete(request, uuid):
    product = get_object_or_404(Product, uuid=uuid)
    if request.method == 'POST':
        try:
            product_name = product.name
            product.delete()
            messages.warning(request, f'Produk "{product_name}" berhasil dihapus.')
        except ProtectedError:
            messages.error(request, 'Produk tidak bisa dihapus karena sudah dipakai transaksi.')
    else:
        messages.info(request, 'Penghapusan dibatalkan.')
    return redirect('product_list')


@role_required(*MANAGEMENT_ROLES, perm='manage_purchases')
def purchase_page(request):
    query = request.GET.get('q', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    txs = InventoryTransaction.objects.select_related('supplier').filter(
        tx_type=InventoryTransaction.TYPE_PURCHASE
    ).order_by('-tx_date', '-created_at')
    if query:
        txs = txs.filter(Q(tx_number__icontains=query) | Q(supplier__name__icontains=query))
    if date_from:
        txs = txs.filter(tx_date__gte=date_from)
    if date_to:
        txs = txs.filter(tx_date__lte=date_to)
    close_dates = set(
        DailyClosing.objects.filter(
            is_locked=True,
            close_date__in=txs.values_list('tx_date', flat=True),
        ).values_list('close_date', flat=True)
    )
    page_obj = Paginator(txs, 10).get_page(request.GET.get('page'))
    for tx in page_obj:
        tx.can_modify = tx.tx_date not in close_dates
    return render(
        request,
        'inventory/purchase_list.html',
        {'page_obj': page_obj, 'query': query, 'date_from': date_from, 'date_to': date_to},
    )


def _extract_purchase_items(request):
    product_ids = request.POST.getlist('product_id[]')
    qtys = request.POST.getlist('qty[]')
    costs = request.POST.getlist('unit_cost[]')
    items = []
    size = max(len(product_ids), len(qtys), len(costs))
    for i in range(size):
        pid = product_ids[i].strip() if i < len(product_ids) else ''
        q = qtys[i].strip() if i < len(qtys) else ''
        c = costs[i].strip() if i < len(costs) else ''
        if not any([pid, q, c]):
            continue
        if not all([pid, q, c]):
            raise ValidationError('Setiap baris item pembelian wajib diisi lengkap.')
        product = Product.objects.get(id=pid)
        items.append({'product': product, 'qty': int(q), 'unit_cost': Decimal(c)})
    return items


def _build_purchase_initial_rows_from_request(request):
    product_ids = request.POST.getlist('product_id[]')
    qtys = request.POST.getlist('qty[]')
    costs = request.POST.getlist('unit_cost[]')
    rows = []
    size = max(len(product_ids), len(qtys), len(costs))
    for i in range(size):
        pid = product_ids[i].strip() if i < len(product_ids) else ''
        q = qtys[i].strip() if i < len(qtys) else ''
        c = costs[i].strip() if i < len(costs) else ''
        if not any([pid, q, c]):
            continue
        rows.append({'product_id': pid, 'qty': q or '1', 'unit_cost': c or ''})
    return rows


@role_required(*MANAGEMENT_ROLES, perm='manage_purchases')
def purchase_create(request):
    products = Product.objects.select_related('unit').order_by('name')
    suppliers = Supplier.objects.filter(is_active=True).order_by('name')
    initial_items_json = _to_json_payload([])
    selected_supplier_id = ''
    if request.method == 'POST':
        initial_items_json = _to_json_payload(_build_purchase_initial_rows_from_request(request))
        selected_supplier_id = request.POST.get('supplier_id', '').strip()
        try:
            supplier_id = request.POST.get('supplier_id', '').strip()
            if not supplier_id:
                raise ValidationError('Supplier wajib dipilih sebelum menyimpan transaksi kulakan.')
            supplier = Supplier.objects.filter(id=supplier_id, is_active=True).first()
            if not supplier:
                raise ValidationError('Supplier tidak ditemukan atau tidak aktif.')
            tx_date_raw = request.POST.get('tx_date', '').strip()
            tx_date = date.fromisoformat(tx_date_raw) if tx_date_raw else date.today()
            note = request.POST.get('note', '').strip()
            items = _extract_purchase_items(request)
            create_purchase_transaction(
                supplier=supplier,
                tx_date=tx_date,
                items=items,
                user=request.user,
                note=note,
            )
            messages.success(request, 'Transaksi pembelian berhasil disimpan.')
            return redirect('purchase_page')
        except Exception as exc:
            messages.error(request, _exc_message(exc))
    suppliers_data = [
        {
            'id': str(s.id),
            'code': s.code or '',
            'name': s.name,
            'phone': s.phone or '',
            'email': s.email or '',
            'address': s.address or '',
            'city': s.city or '',
        }
        for s in suppliers
    ]
    products_data = [
        {
            'id': str(p.id),
            'name': p.name,
            'sku': p.sku,
            'barcode': p.barcode or '',
            'unit': p.unit.name if p.unit else '-',
        }
        for p in products
    ]
    return render(
        request,
        'inventory/purchase_create.html',
        {
            'products': products,
            'suppliers': suppliers,
            'suppliers_json': _to_json_payload(suppliers_data),
            'products_json': _to_json_payload(products_data),
            'initial_items_json': initial_items_json,
            'selected_supplier_id': selected_supplier_id,
        },
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_purchases')
def purchase_detail(request, uuid):
    tx = get_object_or_404(
        InventoryTransaction.objects.select_related('supplier').prefetch_related('items__product'),
        uuid=uuid,
        tx_type=InventoryTransaction.TYPE_PURCHASE,
    )
    can_modify = not DailyClosing.objects.filter(close_date=tx.tx_date, is_locked=True).exists()
    return render(request, 'inventory/purchase_detail.html', {'tx': tx, 'can_modify': can_modify})


@role_required(*MANAGEMENT_ROLES, perm='manage_purchases')
def purchase_edit(request, uuid):
    tx = get_object_or_404(
        InventoryTransaction.objects.select_related('supplier').prefetch_related('items__product'),
        uuid=uuid,
        tx_type=InventoryTransaction.TYPE_PURCHASE,
    )
    if DailyClosing.objects.filter(close_date=tx.tx_date, is_locked=True).exists():
        messages.error(request, 'Transaksi pembelian ini tidak bisa diedit karena tanggalnya sudah tutup harian.')
        return redirect('purchase_detail', uuid=tx.uuid)
    products = Product.objects.select_related('unit').order_by('name')
    suppliers = Supplier.objects.filter(is_active=True).order_by('name')
    next_url = _get_safe_next_url(request)
    back_url = next_url or f'/inventory/purchases/{tx.uuid}/'
    selected_supplier_id = str(tx.supplier_id) if tx.supplier_id else ''
    initial_items_json = _to_json_payload(
        [
            {
                'product_id': str(it.product_id),
                'qty': it.qty,
                'unit_cost': str(it.unit_cost),
            }
            for it in tx.items.all()
        ]
    )
    if request.method == 'POST':
        selected_supplier_id = request.POST.get('supplier_id', '').strip()
        initial_items_json = _to_json_payload(_build_purchase_initial_rows_from_request(request))
        try:
            supplier_id = request.POST.get('supplier_id', '').strip()
            if not supplier_id:
                raise ValidationError('Supplier wajib dipilih sebelum menyimpan transaksi kulakan.')
            supplier = Supplier.objects.filter(id=supplier_id, is_active=True).first()
            if not supplier:
                raise ValidationError('Supplier tidak ditemukan atau tidak aktif.')
            tx_date_raw = request.POST.get('tx_date', '').strip()
            tx_date = date.fromisoformat(tx_date_raw) if tx_date_raw else tx.tx_date
            note = request.POST.get('note', '').strip()
            items = _extract_purchase_items(request)
            edit_purchase_transaction(
                tx=tx,
                supplier=supplier,
                tx_date=tx_date,
                items=items,
                user=request.user,
                note=note,
            )
            messages.success(request, 'Transaksi pembelian berhasil diperbarui.')
            return redirect('purchase_page')
        except Exception as exc:
            messages.error(request, _exc_message(exc))
    return render(
        request,
        'inventory/purchase_edit.html',
        {
            'tx': tx,
            'products': products,
            'suppliers': suppliers,
            'next_url': next_url,
            'back_url': back_url,
            'selected_supplier_id': selected_supplier_id,
            'suppliers_json': _to_json_payload(
                [
                    {
                        'id': str(s.id),
                        'code': s.code or '',
                        'name': s.name,
                        'phone': s.phone or '',
                        'email': s.email or '',
                        'address': s.address or '',
                        'city': s.city or '',
                    }
                    for s in suppliers
                ]
            ),
            'products_json': _to_json_payload(
                [
                    {
                        'id': str(p.id),
                        'name': p.name,
                        'sku': p.sku,
                        'barcode': p.barcode or '',
                        'unit': p.unit.name if p.unit else '-',
                    }
                    for p in products
                ]
            ),
            'initial_items_json': initial_items_json,
        },
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_purchases')
def purchase_delete(request, uuid):
    tx = get_object_or_404(
        InventoryTransaction.objects.select_related('supplier').prefetch_related('items__product'),
        uuid=uuid,
        tx_type=InventoryTransaction.TYPE_PURCHASE,
    )
    if DailyClosing.objects.filter(close_date=tx.tx_date, is_locked=True).exists():
        messages.error(request, 'Transaksi pembelian ini tidak bisa dihapus karena tanggalnya sudah tutup harian.')
        return redirect('purchase_detail', uuid=tx.uuid)
    if request.method == 'POST':
        try:
            delete_purchase_transaction(tx)
            messages.warning(request, 'Transaksi pembelian berhasil dihapus.')
        except Exception as exc:
            messages.error(request, _exc_message(exc))
    return redirect('purchase_page')


@role_required(*MANAGEMENT_ROLES, perm='manage_purchases')
def internal_used_page(request):
    products = Product.objects.order_by('name')
    if request.method == 'POST':
        try:
            product = Product.objects.get(id=request.POST.get('product_id'))
            qty = int(request.POST.get('qty', '0'))
            note = request.POST.get('note', '').strip()
            post_internal_used(product=product, qty=qty, user=request.user, note=note)
            messages.success(request, 'Transaksi internal used berhasil diposting.')
            return redirect('internal_used_page')
        except Exception as exc:
            messages.error(request, _exc_message(exc))
    return render(request, 'inventory/internal_used_page.html', {'products': products})


@role_required(*MANAGEMENT_ROLES, perm='perform_stock_opname')
def stock_opname_page(request):
    products = Product.objects.order_by('name')
    if request.method == 'POST':
        try:
            product = Product.objects.get(id=request.POST.get('product_id'))
            actual_stock = int(request.POST.get('actual_stock', '0'))
            note = request.POST.get('note', '').strip()
            post_stock_opname(product=product, actual_stock=actual_stock, user=request.user, note=note)
            messages.success(request, 'Transaksi stock opname berhasil diposting.')
            return redirect('stock_opname_page')
        except Exception as exc:
            messages.error(request, _exc_message(exc))
    return render(request, 'inventory/stock_opname_page.html', {'products': products})


@role_required(Role.ADMIN_TOKO, perm='perform_daily_closing')
def daily_closing_page(request):
    if request.method == 'POST':
        try:
            action = request.POST.get('action', 'close')
            if action == 'reopen':
                reopen_last_closing(user=request.user)
                messages.warning(request, 'Closing tanggal terakhir berhasil dibuka kembali.')
                return redirect('daily_closing_page')

            close_date_raw = request.POST.get('close_date', '')
            close_date = date.fromisoformat(close_date_raw) if close_date_raw else date.today()
            note = request.POST.get('note', '').strip()
            close_daily(closing_date=close_date, user=request.user, note=note)
            messages.success(request, f'Tutup harian {close_date} berhasil.')
            return redirect('daily_closing_page')
        except Exception as exc:
            messages.error(request, _exc_message(exc))
    latest_closing = DailyClosing.objects.order_by('-close_date').first()
    next_closing_date = date.today()
    if latest_closing:
        next_closing_date = latest_closing.close_date + timedelta(days=1)
    return render(
        request,
        'inventory/daily_closing_page.html',
        {'latest_closing': latest_closing, 'next_closing_date': next_closing_date},
    )


@role_required(Role.ADMIN_TOKO, perm='perform_daily_closing')
def daily_closing_report(request):
    close_date = request.GET.get('close_date', '').strip()
    closing = None
    if close_date:
        closing = DailyClosing.objects.filter(close_date=close_date).first()
    if not closing:
        closing = DailyClosing.objects.order_by('-close_date').first()

    product_rows = []
    member_rows = []
    product_mismatch_count = 0
    member_mismatch_count = 0

    if closing:
        for snap in closing.product_snapshots.select_related('product').order_by('product__name'):
            current_stock = snap.product.stock
            diff = current_stock - snap.closing_stock
            is_match = diff == 0
            if not is_match:
                product_mismatch_count += 1
            product_rows.append(
                {
                    'product': snap.product,
                    'opening': snap.opening_stock,
                    'mut_in': snap.mutation_in,
                    'mut_out': snap.mutation_out,
                    'closing': snap.closing_stock,
                    'actual': current_stock,
                    'diff': diff,
                    'is_match': is_match,
                }
            )
        for snap in closing.member_snapshots.select_related('member').order_by('member__full_name'):
            wallet = getattr(snap.member, 'wallet', None)
            current_balance = wallet.balance if wallet else Decimal('0.00')
            diff = current_balance - snap.closing_balance
            is_match = diff == 0
            if not is_match:
                member_mismatch_count += 1
            member_rows.append(
                {
                    'member': snap.member,
                    'opening': snap.opening_balance,
                    'mut_in': snap.mutation_in,
                    'mut_out': snap.mutation_out,
                    'closing': snap.closing_balance,
                    'actual': current_balance,
                    'diff': diff,
                    'is_match': is_match,
                }
            )

    close_dates = DailyClosing.objects.order_by('-close_date').values_list('close_date', flat=True)
    return render(
        request,
        'inventory/daily_closing_report.html',
        {
            'closing': closing,
            'close_dates': close_dates,
            'selected_date': close_date,
            'product_rows': product_rows,
            'member_rows': member_rows,
            'product_mismatch_count': product_mismatch_count,
            'member_mismatch_count': member_mismatch_count,
        },
    )


@role_required(*STAFF_ROLES, perm='view_inventory')
def stock_card_report(request):
    products = Product.objects.select_related('unit').order_by('name')
    product_id = request.GET.get('product_id', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    ledgers = []
    selected_product = None
    if product_id:
        selected_product = Product.objects.filter(id=product_id).first()
        if selected_product:
            ledgers = selected_product.stock_ledgers.select_related('tx').order_by('tx_date', 'created_at')
            if date_from:
                ledgers = ledgers.filter(tx_date__gte=date_from)
            if date_to:
                ledgers = ledgers.filter(tx_date__lte=date_to)
    return render(
        request,
        'inventory/stock_card_report.html',
        {
            'products': products,
            'products_json': _to_json_payload(
                [
                    {
                        'id': str(p.id),
                        'name': p.name,
                        'sku': p.sku,
                        'barcode': p.barcode or '',
                        'unit': p.unit.name if p.unit else '-',
                        'stock': p.stock,
                    }
                    for p in products
                ]
            ),
            'selected_product': selected_product,
            'ledgers': ledgers,
            'date_from': date_from,
            'date_to': date_to,
        },
    )


@role_required(*MANAGEMENT_ROLES, perm='view_inventory')
def reorder_alert_page(request):
    return render(request, 'inventory/reorder_alert_page.html', {'products': low_stock_products()})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def category_list(request):
    query = request.GET.get('q', '').strip()
    categories = Category.objects.order_by('name')
    if query:
        categories = categories.filter(name__icontains=query)

    paginator = Paginator(categories, 10)
    page_obj = paginator.get_page(request.GET.get('page'))
    return render(
        request,
        'inventory/category_list.html',
        {'page_obj': page_obj, 'query': query},
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def category_create(request):
    error_message = ''
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        if name:
            Category.objects.create(name=name)
            messages.success(request, 'Kategori berhasil ditambahkan.')
            return redirect('category_list')
        error_message = 'Nama kategori wajib diisi.'
        messages.error(request, error_message)
    return render(
        request,
        'inventory/category_create.html',
        {'error_message': error_message},
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def category_detail(request, uuid):
    category = get_object_or_404(Category, uuid=uuid)
    return render(request, 'inventory/category_detail.html', {'category': category})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def category_edit(request, uuid):
    category = get_object_or_404(Category, uuid=uuid)
    error_message = ''
    next_url = _get_safe_next_url(request)
    back_url = next_url or f"/inventory/categories/{category.uuid}/"

    if request.method == 'POST':
        category.name = request.POST.get('name', '').strip()
        if category.name:
            category.save()
            messages.success(request, 'Kategori berhasil diperbarui.')
            return redirect('category_list')
        error_message = 'Nama kategori wajib diisi.'
        messages.error(request, error_message)
    return render(
        request,
        'inventory/category_edit.html',
        {'category': category, 'error_message': error_message, 'back_url': back_url, 'next_url': next_url},
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def category_delete(request, uuid):
    category = get_object_or_404(Category, uuid=uuid)
    if request.method == 'POST':
        try:
            category_name = category.name
            category.delete()
            messages.warning(request, f'Kategori "{category_name}" berhasil dihapus.')
        except ProtectedError:
            messages.error(request, 'Kategori tidak bisa dihapus karena sudah dipakai produk/transaksi.')
    else:
        messages.info(request, 'Penghapusan dibatalkan.')
    return redirect('category_list')


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def unit_list(request):
    query = request.GET.get('q', '').strip()
    units = Unit.objects.order_by('name')
    if query:
        units = units.filter(Q(name__icontains=query) | Q(code__icontains=query))

    page_obj = Paginator(units, 10).get_page(request.GET.get('page'))
    return render(request, 'inventory/unit_list.html', {'page_obj': page_obj, 'query': query})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def unit_create(request):
    error_message = ''
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        code = request.POST.get('code', '').strip().upper()
        description = request.POST.get('description', '').strip()
        is_active = request.POST.get('is_active') == 'on'
        if not name or not code:
            error_message = 'Nama dan kode satuan wajib diisi.'
            messages.error(request, error_message)
        elif Unit.objects.filter(code__iexact=code).exists():
            error_message = 'Kode satuan sudah dipakai.'
            messages.error(request, error_message)
        else:
            Unit.objects.create(name=name, code=code, description=description, is_active=is_active)
            messages.success(request, 'Satuan berhasil ditambahkan.')
            return redirect('unit_list')
    return render(request, 'inventory/unit_create.html', {'error_message': error_message})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def unit_detail(request, uuid):
    unit = get_object_or_404(Unit, uuid=uuid)
    return render(request, 'inventory/unit_detail.html', {'unit': unit})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def unit_edit(request, uuid):
    unit = get_object_or_404(Unit, uuid=uuid)
    next_url = _get_safe_next_url(request)
    back_url = next_url or f'/inventory/units/{unit.uuid}/'
    error_message = ''
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        code = request.POST.get('code', '').strip().upper()
        description = request.POST.get('description', '').strip()
        is_active = request.POST.get('is_active') == 'on'
        if not name or not code:
            error_message = 'Nama dan kode satuan wajib diisi.'
            messages.error(request, error_message)
        elif Unit.objects.exclude(id=unit.id).filter(code__iexact=code).exists():
            error_message = 'Kode satuan sudah dipakai.'
            messages.error(request, error_message)
        else:
            unit.name = name
            unit.code = code
            unit.description = description
            unit.is_active = is_active
            unit.save()
            messages.success(request, 'Satuan berhasil diperbarui.')
            return redirect('unit_list')
    return render(
        request,
        'inventory/unit_edit.html',
        {'unit': unit, 'next_url': next_url, 'back_url': back_url, 'error_message': error_message},
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def unit_delete(request, uuid):
    unit = get_object_or_404(Unit, uuid=uuid)
    if request.method == 'POST':
        try:
            name = unit.name
            unit.delete()
            messages.warning(request, f'Satuan "{name}" berhasil dihapus.')
        except Exception:
            messages.error(request, 'Satuan tidak bisa dihapus karena sudah dipakai produk.')
    return redirect('unit_list')


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def supplier_list(request):
    query = request.GET.get('q', '').strip()
    suppliers = Supplier.objects.order_by('name')
    if query:
        suppliers = suppliers.filter(
            Q(code__icontains=query)
            | Q(name__icontains=query)
            | Q(contact_name__icontains=query)
            | Q(phone__icontains=query)
            | Q(email__icontains=query)
            | Q(city__icontains=query)
        )
    page_obj = Paginator(suppliers, 10).get_page(request.GET.get('page'))
    return render(request, 'inventory/supplier_list.html', {'page_obj': page_obj, 'query': query})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def supplier_create(request):
    error_message = ''
    if request.method == 'POST':
        code = request.POST.get('code', '').strip().upper()
        name = request.POST.get('name', '').strip()
        contact_name = request.POST.get('contact_name', '').strip()
        phone = request.POST.get('phone', '').strip()
        email = request.POST.get('email', '').strip()
        address = request.POST.get('address', '').strip()
        city = request.POST.get('city', '').strip()
        is_active = request.POST.get('is_active') == 'on'
        if not code or not name:
            error_message = 'Kode dan nama supplier wajib diisi.'
            messages.error(request, error_message)
        elif Supplier.objects.filter(code__iexact=code).exists():
            error_message = 'Kode supplier sudah dipakai.'
            messages.error(request, error_message)
        elif Supplier.objects.filter(name__iexact=name).exists():
            error_message = 'Nama supplier sudah dipakai.'
            messages.error(request, error_message)
        else:
            Supplier.objects.create(
                code=code,
                name=name,
                contact_name=contact_name,
                phone=phone,
                email=email,
                address=address,
                city=city,
                is_active=is_active,
                created_by=request.user,
                updated_by=request.user,
            )
            messages.success(request, 'Supplier berhasil ditambahkan.')
            return redirect('supplier_list')
    return render(request, 'inventory/supplier_create.html', {'error_message': error_message})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def supplier_detail(request, uuid):
    supplier = get_object_or_404(Supplier, uuid=uuid)
    return render(request, 'inventory/supplier_detail.html', {'supplier': supplier})


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def supplier_edit(request, uuid):
    supplier = get_object_or_404(Supplier, uuid=uuid)
    next_url = _get_safe_next_url(request)
    back_url = next_url or f'/inventory/suppliers/{supplier.uuid}/'
    error_message = ''
    if request.method == 'POST':
        code = request.POST.get('code', '').strip().upper()
        name = request.POST.get('name', '').strip()
        contact_name = request.POST.get('contact_name', '').strip()
        phone = request.POST.get('phone', '').strip()
        email = request.POST.get('email', '').strip()
        address = request.POST.get('address', '').strip()
        city = request.POST.get('city', '').strip()
        is_active = request.POST.get('is_active') == 'on'
        if not code or not name:
            error_message = 'Kode dan nama supplier wajib diisi.'
            messages.error(request, error_message)
        elif Supplier.objects.exclude(id=supplier.id).filter(code__iexact=code).exists():
            error_message = 'Kode supplier sudah dipakai.'
            messages.error(request, error_message)
        elif Supplier.objects.exclude(id=supplier.id).filter(name__iexact=name).exists():
            error_message = 'Nama supplier sudah dipakai.'
            messages.error(request, error_message)
        else:
            supplier.code = code
            supplier.name = name
            supplier.contact_name = contact_name
            supplier.phone = phone
            supplier.email = email
            supplier.address = address
            supplier.city = city
            supplier.is_active = is_active
            supplier.updated_by = request.user
            supplier.save()
            messages.success(request, 'Supplier berhasil diperbarui.')
            return redirect('supplier_list')
    return render(
        request,
        'inventory/supplier_edit.html',
        {'supplier': supplier, 'next_url': next_url, 'back_url': back_url, 'error_message': error_message},
    )


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def supplier_delete(request, uuid):
    supplier = get_object_or_404(Supplier, uuid=uuid)
    if request.method == 'POST':
        try:
            supplier_name = supplier.name
            supplier.delete()
            messages.warning(request, f'Supplier "{supplier_name}" berhasil dihapus.')
        except Exception:
            messages.error(request, 'Supplier tidak bisa dihapus karena sudah dipakai transaksi.')
    return redirect('supplier_list')


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def product_import_template_excel(request):
    """
    Download template Excel (.xlsx) resmi untuk import data master barang.
    Sheet 1: Template Import (kolom sku, barcode, nama_barang, kategori, satuan, harga_beli, harga_jual, min_stok, stok_awal)
    Sheet 2: Referensi Master Data (daftar Kategori dan Satuan aktif di database)
    """
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse

    wb = openpyxl.Workbook()

    # Sheet 1: Template Import
    ws1 = wb.active
    ws1.title = "Template Import"

    headers = [
        ("sku", 18),
        ("barcode", 20),
        ("nama_barang", 35),
        ("kategori", 20),
        ("satuan", 15),
        ("harga_beli", 16),
        ("harga_jual", 16),
        ("min_stok", 14),
        ("stok_awal", 14),
        ("bisa_desimal", 16),
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    # Write headers
    for col_idx, (col_id, col_width) in enumerate(headers, start=1):
        cell = ws1.cell(row=1, column=col_idx, value=col_id)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = col_width

    # Sample rows to demonstrate formatting
    sample_rows = [
        ["BRG-001", "8991234567890", "Beras Ramos 5kg", "Sembako", "PCS", 60000, 68000, 10, 50, "TIDAK"],
        ["BRG-002", "", "Gula Pasir (Curah)", "Sembako", "KG", 15000, 17500, 20, 100, "YA"],
    ]
    sample_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    for row_idx, srow in enumerate(sample_rows, start=2):
        for col_idx, val in enumerate(srow, start=1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.fill = sample_fill
            if col_idx in [6, 7, 8, 9]:
                cell.alignment = right_align
            elif col_idx in [1, 2, 5, 10]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Sheet 2: Referensi Master Data
    ws2 = wb.create_sheet(title="Referensi Master Data")
    ref_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ref_header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")

    # Categories in Col A
    cat_header = ws2.cell(row=1, column=1, value="Kategori Terdaftar")
    cat_header.font = ref_header_font
    cat_header.fill = ref_header_fill
    cat_header.alignment = center_align
    cat_header.border = thin_border
    ws2.column_dimensions['A'].width = 28

    categories = Category.objects.all().order_by('name')
    for idx, c in enumerate(categories, start=2):
        cell = ws2.cell(row=idx, column=1, value=c.name)
        cell.border = thin_border

    # Units in Col C, D
    unit_code_header = ws2.cell(row=1, column=3, value="Kode Satuan")
    unit_code_header.font = ref_header_font
    unit_code_header.fill = ref_header_fill
    unit_code_header.alignment = center_align
    unit_code_header.border = thin_border
    ws2.column_dimensions['C'].width = 18

    unit_name_header = ws2.cell(row=1, column=4, value="Nama Satuan")
    unit_name_header.font = ref_header_font
    unit_name_header.fill = ref_header_fill
    unit_name_header.alignment = center_align
    unit_name_header.border = thin_border
    ws2.column_dimensions['D'].width = 25

    active_units = Unit.objects.filter(is_active=True).order_by('name')
    for idx, u in enumerate(active_units, start=2):
        cell_code = ws2.cell(row=idx, column=3, value=u.code)
        cell_code.border = thin_border
        cell_code.alignment = center_align
        cell_name = ws2.cell(row=idx, column=4, value=u.name)
        cell_name.border = thin_border

    ws2.column_dimensions['B'].width = 5

    # Instructions note on Sheet 2
    note_cell = ws2.cell(row=1, column=6, value="PERHATIAN:")
    note_cell.font = Font(name="Calibri", size=11, bold=True, color="DC2626")
    ws2.cell(row=2, column=6, value="1. Sistem TIDAK membuat Kategori dan Satuan baru secara otomatis.")
    ws2.cell(row=3, column=6, value="2. Nama kategori dan satuan di sheet 'Template Import' harus persis sama dengan daftar master di atas.")
    ws2.cell(row=4, column=6, value="3. Jika ada kategori/satuan yang tidak cocok atau SKU duplikat, seluruh data import akan ditolak.")
    ws2.column_dimensions['F'].width = 70

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="template_import_produk.xlsx"'
    return response


@role_required(*MANAGEMENT_ROLES, perm='manage_products')
def product_import_excel(request):
    """
    Import data master barang dari file Excel (.xlsx) dengan validasi ketat:
    - Tidak ada auto-create kategori atau satuan.
    - Menolak jika ada SKU yang sudah ada di database atau duplikat di file.
    - All-or-nothing: Jika ada 1 baris error, seluruh data tidak bisa disimpan.
    """
    import openpyxl

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'cancel':
            request.session.pop('product_import_batch', None)
            messages.info(request, 'Proses import dibatalkan.')
            return redirect('product_import_excel')

        elif action == 'confirm':
            batch = request.session.get('product_import_batch')
            if not batch or not batch.get('can_confirm') or batch.get('total_errors', 1) > 0:
                messages.error(request, 'Data import tidak valid, memiliki error, atau sesi telah kedaluwarsa. Silakan unggah ulang.')
                return redirect('product_import_excel')

            validated_rows = batch.get('rows', [])
            try:
                created_count = import_products_batch(validated_rows, request.user)
                request.session.pop('product_import_batch', None)
                messages.success(request, f'Berhasil mengimpor {created_count} produk ke data master.')
                return redirect('product_list')
            except Exception as e:
                messages.error(request, f'Gagal menyimpan data import: {_exc_message(e)}')
                return render(request, 'inventory/product_import.html', {
                    'preview_mode': True,
                    'rows': validated_rows,
                    'total_rows': batch.get('total_rows', 0),
                    'total_errors': batch.get('total_errors', 0),
                    'total_valid': batch.get('total_valid', 0),
                    'can_confirm': False,
                    'system_error': _exc_message(e),
                })

        elif action == 'preview':
            uploaded_file = request.FILES.get('file')
            if not uploaded_file:
                messages.error(request, 'Silakan pilih file Excel (.xlsx) terlebih dahulu.')
                return render(request, 'inventory/product_import.html')

            if not uploaded_file.name.lower().endswith(('.xlsx', '.xlsm')):
                messages.error(request, 'Format file tidak didukung. Harap unggah file berformat .xlsx.')
                return render(request, 'inventory/product_import.html')

            try:
                wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            except Exception as e:
                messages.error(request, f'Gagal membaca file Excel: {str(e)}')
                return render(request, 'inventory/product_import.html')

            ws = wb['Template Import'] if 'Template Import' in wb.sheetnames else wb.active

            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                messages.error(request, 'File Excel kosong.')
                return render(request, 'inventory/product_import.html')

            if not header_row:
                messages.error(request, 'Baris header tidak ditemukan pada file Excel.')
                return render(request, 'inventory/product_import.html')

            # Map header columns
            col_map = {}
            for idx, cell_value in enumerate(header_row):
                if cell_value is not None:
                    norm_key = str(cell_value).strip().lower().replace(' ', '_').replace('*', '').strip('_')
                    col_map[norm_key] = idx

            # Mandatory headers
            required_cols = ['sku', 'nama_barang', 'kategori', 'satuan', 'harga_jual']
            if 'nama' in col_map and 'nama_barang' not in col_map:
                col_map['nama_barang'] = col_map['nama']

            missing_headers = [c for c in required_cols if c not in col_map]
            if missing_headers:
                messages.error(
                    request,
                    f"Kolom wajib tidak ditemukan di file Excel: {', '.join(missing_headers)}. Silakan gunakan template resmi."
                )
                return render(request, 'inventory/product_import.html')

            # Pre-load master categories and units for fast and case-insensitive matching
            categories = {c.name.strip().lower(): c for c in Category.objects.all()}
            units = {}
            for u in Unit.objects.filter(is_active=True):
                units[u.code.strip().lower()] = u
                units[u.name.strip().lower()] = u

            # Existing SKUs and barcodes in DB
            existing_skus_db = {s.lower() for s in Product.objects.values_list('sku', flat=True)}
            existing_barcodes_db = {
                b.lower()
                for b in Product.objects.exclude(barcode__isnull=True).exclude(barcode='').values_list('barcode', flat=True)
            }

            seen_skus_file = {}
            seen_barcodes_file = {}

            parsed_rows = []
            total_errors = 0
            total_valid = 0

            for row_idx, row_values in enumerate(rows_iter, start=2):
                if not any(v is not None and str(v).strip() != '' for v in row_values):
                    continue  # skip completely blank line

                def get_val(col_name, default=''):
                    idx = col_map.get(col_name)
                    if idx is not None and idx < len(row_values):
                        val = row_values[idx]
                        return str(val).strip() if val is not None else default
                    return default

                raw_sku = get_val('sku')
                raw_barcode = get_val('barcode')
                raw_nama = get_val('nama_barang')
                raw_kategori = get_val('kategori')
                raw_satuan = get_val('satuan')
                raw_harga_beli = get_val('harga_beli', '0')
                raw_harga_jual = get_val('harga_jual')
                raw_min_stok = get_val('min_stok', '0')
                raw_stok_awal = get_val('stok_awal', '0')

                row_errors = []

                # 1. Mandatory checks
                if not raw_sku:
                    row_errors.append("SKU wajib diisi.")
                if not raw_nama:
                    row_errors.append("Nama barang wajib diisi.")
                if not raw_kategori:
                    row_errors.append("Kategori wajib diisi.")
                if not raw_satuan:
                    row_errors.append("Satuan wajib diisi.")
                if not raw_harga_jual:
                    row_errors.append("Harga jual wajib diisi.")

                # 2. SKU checks
                if raw_sku:
                    sku_lower = raw_sku.lower()
                    if sku_lower in seen_skus_file:
                        row_errors.append(f"SKU '{raw_sku}' duplikat dengan baris {seen_skus_file[sku_lower]} di file Excel.")
                    else:
                        seen_skus_file[sku_lower] = row_idx

                    if sku_lower in existing_skus_db:
                        row_errors.append(f"SKU '{raw_sku}' sudah terdaftar di database.")

                # 3. Barcode checks
                if raw_barcode:
                    bc_lower = raw_barcode.lower()
                    if bc_lower in seen_barcodes_file:
                        row_errors.append(f"Barcode '{raw_barcode}' duplikat dengan baris {seen_barcodes_file[bc_lower]} di file Excel.")
                    else:
                        seen_barcodes_file[bc_lower] = row_idx

                    if bc_lower in existing_barcodes_db:
                        row_errors.append(f"Barcode '{raw_barcode}' sudah digunakan produk lain di database.")

                # 4. Kategori master check (NO auto-create!)
                cat_obj = None
                if raw_kategori:
                    cat_obj = categories.get(raw_kategori.lower())
                    if not cat_obj:
                        row_errors.append(f"Kategori '{raw_kategori}' tidak ditemukan di data master.")

                # 5. Satuan master check (NO auto-create!)
                unit_obj = None
                if raw_satuan:
                    unit_obj = units.get(raw_satuan.lower())
                    if not unit_obj:
                        row_errors.append(f"Satuan '{raw_satuan}' tidak ditemukan atau nonaktif di data master.")

                # 6. Numeric checks
                parsed_harga_jual = Decimal('0')
                if raw_harga_jual:
                    try:
                        clean_hj = raw_harga_jual.replace(',', '').replace(' ', '')
                        parsed_harga_jual = Decimal(clean_hj)
                        if parsed_harga_jual <= 0:
                            row_errors.append("Harga jual harus lebih besar dari 0.")
                    except (InvalidOperation, ValueError):
                        row_errors.append(f"Format harga jual tidak valid: '{raw_harga_jual}'.")

                parsed_harga_beli = Decimal('0')
                if raw_harga_beli:
                    try:
                        clean_hb = raw_harga_beli.replace(',', '').replace(' ', '')
                        parsed_harga_beli = Decimal(clean_hb)
                        if parsed_harga_beli < 0:
                            row_errors.append("Harga beli tidak boleh negatif.")
                    except (InvalidOperation, ValueError):
                        row_errors.append(f"Format harga beli tidak valid: '{raw_harga_beli}'.")

                parsed_min_stok = Decimal('0')
                if raw_min_stok:
                    try:
                        clean_ms = raw_min_stok.replace(',', '').replace(' ', '')
                        parsed_min_stok = Decimal(clean_ms)
                        if parsed_min_stok < Decimal('0'):
                            row_errors.append("Min stok tidak boleh negatif.")
                    except (InvalidOperation, ValueError):
                        row_errors.append(f"Format min stok tidak valid: '{raw_min_stok}'.")

                parsed_stok_awal = Decimal('0')
                if raw_stok_awal:
                    try:
                        clean_sa = raw_stok_awal.replace(',', '').replace(' ', '')
                        parsed_stok_awal = Decimal(clean_sa)
                        if parsed_stok_awal < Decimal('0'):
                            row_errors.append("Stok awal tidak boleh negatif.")
                    except (InvalidOperation, ValueError):
                        row_errors.append(f"Format stok awal tidak valid: '{raw_stok_awal}'.")

                raw_bisa_desimal = get_val('bisa_desimal') or get_val('curah') or get_val('desimal')
                is_decimal = raw_bisa_desimal.lower() in ['ya', 'y', 'true', '1', 'yes', 'curah']

                is_valid = len(row_errors) == 0
                if is_valid:
                    total_valid += 1
                else:
                    total_errors += 1

                parsed_rows.append({
                    'row_idx': row_idx,
                    'sku': raw_sku,
                    'barcode': raw_barcode,
                    'nama_barang': raw_nama,
                    'kategori': raw_kategori,
                    'category_id': cat_obj.id if cat_obj else None,
                    'satuan': raw_satuan,
                    'unit_id': unit_obj.id if unit_obj else None,
                    'unit_display': f"{unit_obj.name} ({unit_obj.code})" if unit_obj else raw_satuan,
                    'harga_beli': str(parsed_harga_beli),
                    'harga_jual': str(parsed_harga_jual),
                    'min_stok': str(parsed_min_stok),
                    'stok_awal': str(parsed_stok_awal),
                    'allow_decimal_qty': is_decimal,
                    'is_valid': is_valid,
                    'errors': row_errors,
                })

            total_rows = len(parsed_rows)
            if total_rows == 0:
                messages.warning(request, 'Tidak ada baris data barang yang ditemukan di file Excel.')
                return render(request, 'inventory/product_import.html')

            can_confirm = (total_rows > 0 and total_errors == 0)

            # Store in session for confirmation
            request.session['product_import_batch'] = {
                'rows': parsed_rows,
                'can_confirm': can_confirm,
                'total_rows': total_rows,
                'total_errors': total_errors,
                'total_valid': total_valid,
            }

            return render(request, 'inventory/product_import.html', {
                'preview_mode': True,
                'rows': parsed_rows,
                'can_confirm': can_confirm,
                'total_rows': total_rows,
                'total_errors': total_errors,
                'total_valid': total_valid,
            })

    # GET request
    return render(request, 'inventory/product_import.html', {
        'preview_mode': False,
    })
